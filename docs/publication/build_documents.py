"""Create Word copies of the preserved review Markdown using python-docx.

Input is HTML parsed by the bundled marked module. Source content is retained;
Markdown headings, links, code, tables and local images become Word structures.
Run with the Codex bundled Python runtime. QA outputs belong outside the repo.
"""
import argparse
import hashlib
import json
import re
from pathlib import Path
from urllib.parse import quote, unquote, urlparse

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION_START, WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from lxml import html
from PIL import Image

parser = argparse.ArgumentParser()
parser.add_argument('root', type=Path)
parser.add_argument('parsed', type=Path)
parser.add_argument('--only', help='Rebuild one source filename without rewriting other Word files')
args = parser.parse_args()
ROOT = args.root.resolve()
OUT = ROOT / 'docs/documents'
BASE = 'https://github.com/Xenaquas/Noteable/blob/main/'
manifest_path = ROOT / 'docs/publication/document-manifest.json'
records = json.loads(manifest_path.read_text(encoding='utf8')) if args.only and manifest_path.exists() else []

def clean_heading(s):
    return re.sub(r'\s+', ' ', re.sub(r'[^\w\s]', ' ', s).replace('_', ' ')).strip()

def external(href, source):
    if urlparse(href).scheme or href.startswith('#'):
        return href
    target = (source.parent / unquote(href.split('#')[0])).resolve()
    try:
        return BASE + quote(target.relative_to(ROOT).as_posix()) + ('#' + href.split('#', 1)[1] if '#' in href else '')
    except ValueError:
        return href

def link(p, text, href, source, bold=False):
    h = OxmlElement('w:hyperlink')
    if href.startswith('#'):
        h.set(qn('w:anchor'), href[1:])
    else:
        h.set(qn('r:id'), p.part.relate_to(external(href, source), 'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink', is_external=True))
    r = OxmlElement('w:r')
    pr = OxmlElement('w:rPr')
    color = OxmlElement('w:color'); color.set(qn('w:val'), '185A83'); pr.append(color)
    if bold:
        pr.append(OxmlElement('w:b'))
    r.append(pr)
    t = OxmlElement('w:t'); t.text = text; t.set(qn('xml:space'), 'preserve'); r.append(t)
    h.append(r); p._p.append(h)

def picture(p, node, source, maxw=6.8, maxh=7.3):
    src = node.get('src', '')
    if urlparse(src).scheme:
        link(p, node.get('alt') or src, src, source)
        return
    image = (source.parent / unquote(src)).resolve()
    if not image.exists():
        p.add_run(node.get('alt') or src)
        return
    with Image.open(image) as im:
        width, height = im.size
    w = min(maxw, maxh * width / height)
    shape = p.add_run().add_picture(str(image), width=Inches(w))
    shape._inline.docPr.set('descr', node.get('alt') or image.name)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER

def inline(p, element, source, width=6.8, bold=False, italic=False, code=False):
    def add(text):
        if not text:
            return
        r = p.add_run(text)
        r.bold = bold; r.italic = italic
        if code:
            r.font.name = 'Consolas'; r.font.size = Pt(9)
    add(element.text)
    for ch in element:
        tag = ch.tag.lower() if isinstance(ch.tag, str) else ''
        if tag == 'img':
            picture(p, ch, source, width, 5.8 if width < 4 else 7.3)
        elif tag == 'a' and not ch.xpath('.//img'):
            link(p, ''.join(ch.itertext()), ch.get('href', ''), source, bold)
        elif tag == 'br':
            p.add_run().add_break()
        else:
            inline(p, ch, source, width, bold or tag in ('strong', 'b'), italic or tag in ('em', 'i'), code or tag == 'code')
        add(ch.tail)

def set_table(table, widths, header=True):
    table.autofit = False
    tblpr = table._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for name in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        e = OxmlElement('w:' + name); e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), '4'); e.set(qn('w:color'), 'D9D9D9'); borders.append(e)
    tblpr.append(borders)
    for col, width in zip(table.columns, widths):
        col.width = Inches(width)
    for i, row in enumerate(table.rows):
        # Keep each comparable record together and repeat the table header.
        row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
        if i == 0 and header:
            row._tr.get_or_add_trPr().append(OxmlElement('w:tblHeader'))
        for j, cell in enumerate(row.cells):
            cell.width = Inches(widths[j]); cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr = cell._tc.get_or_add_tcPr()
            margins = OxmlElement('w:tcMar')
            for side in ('top', 'left', 'bottom', 'right'):
                e = OxmlElement('w:' + side); e.set(qn('w:w'), '95'); e.set(qn('w:type'), 'dxa'); margins.append(e)
            pr.append(margins)
            if header:
                shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), '293C4A' if i == 0 else ('F1F5F7' if i % 2 else 'FFFFFF')); pr.append(shd)
            for p in cell.paragraphs:
                if header and i == 0:
                    p.paragraph_format.keep_with_next = True
                p.paragraph_format.space_after = Pt(4)
                p.paragraph_format.line_spacing = 1.06
                if len(''.join(cell._tc.itertext())) < 35:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for r in p.runs:
                    r.font.size = Pt(9.2)
                    if header and i == 0:
                        r.bold = True; r.font.color.rgb = RGBColor(255, 255, 255)

def evidence_gallery(doc, rows, source):
    for start in range(0, len(rows), 2):
        doc.add_page_break()
        doc.add_heading(f'Screenshot evidence {start + 1} to {min(start + 2, len(rows))}', 2)
        table = doc.add_table(rows=1, cols=2)
        for j, row in enumerate(rows[start:start + 2]):
            cols = row.xpath('./td')
            cell = table.cell(0, j)
            img = cols[0].xpath('.//img')[0]
            caption = cell.paragraphs[0]
            caption.add_run(img.get('alt') or Path(img.get('src')).name).bold = True
            p = cell.add_paragraph()
            picture(p, img, source, 2.65, 5.9)
            p = cell.add_paragraph(); inline(p, cols[1], source, 2.9)
            link(cell.add_paragraph(), 'Original full resolution screenshot', img.get('src'), source)
        set_table(table, [3.4, 3.4], header=False)

def table_block(doc, element, source):
    rows = element.xpath('./thead/tr|./tbody/tr|./tr')
    if len(rows) > 20 and element.xpath('.//img'):
        evidence_gallery(doc, rows[1:], source)
        return
    if not rows:
        return
    n = max(len(r.xpath('./td|./th')) for r in rows)
    counts = [0] * n
    for row in rows:
        for j, c in enumerate(row.xpath('./td|./th')):
            counts[j] += min(180, len(''.join(c.itertext())))
    if element.xpath('.//img'):
        widths = [6.8 / n] * n
    else:
        weights = [max(18, v / len(rows)) ** .65 for v in counts]
        widths = [6.8 * x / sum(weights) for x in weights]
    table = doc.add_table(rows=len(rows), cols=n)
    for i, row in enumerate(rows):
        for j, c in enumerate(row.xpath('./td|./th')):
            p = table.cell(i, j).paragraphs[0]
            inline(p, c, source, max(.5, widths[j] - .18))
    set_table(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)

def block(doc, node, source, list_level=0):
    tag = node.tag.lower() if isinstance(node.tag, str) else ''
    if re.fullmatch('h[1-6]', tag):
        title = clean_heading(''.join(node.itertext()))
        is_title = tag == 'h1' and not doc.paragraphs
        p = doc.add_paragraph(title, 'Title' if is_title else f'Heading {max(1, min(int(tag[1]) - 1, 3))}')
        return
    if tag in ('p', 'blockquote'):
        if tag == 'blockquote':
            for ch in node:
                block(doc, ch, source)
        else:
            p = doc.add_paragraph(); inline(p, node, source)
        return
    if tag in ('ul', 'ol'):
        number_id = None
        if tag == 'ol':
            numbering = doc.part.numbering_part.element
            old_id = doc.styles['List Number'].element.pPr.numPr.numId.val
            old_num = numbering.xpath(f'./w:num[@w:numId="{old_id}"]')[0]
            abstract_id = old_num.find(qn('w:abstractNumId')).get(qn('w:val'))
            number_id = max(int(x.get(qn('w:numId'))) for x in numbering.xpath('./w:num')) + 1
            num = OxmlElement('w:num'); num.set(qn('w:numId'), str(number_id))
            a = OxmlElement('w:abstractNumId'); a.set(qn('w:val'), abstract_id); num.append(a)
            override = OxmlElement('w:lvlOverride'); override.set(qn('w:ilvl'), '0')
            start = OxmlElement('w:startOverride'); start.set(qn('w:val'), '1'); override.append(start); num.append(override)
            numbering.append(num)
        for li in node.xpath('./li'):
            p = doc.add_paragraph(style='List Number' if tag == 'ol' else 'List Bullet')
            if number_id is not None:
                numpr = p._p.get_or_add_pPr().get_or_add_numPr()
                numpr.get_or_add_ilvl().val = 0
                numpr.get_or_add_numId().val = number_id
            p.paragraph_format.left_indent = Inches(.22 + list_level * .2)
            # Inline copy without nested lists; recurse for their own Word list items.
            clone = html.fromstring(html.tostring(li))
            for nested in clone.xpath('./ul|./ol'):
                nested.getparent().remove(nested)
            inline(p, clone, source, 6.4)
            for nested in li.xpath('./ul|./ol'):
                block(doc, nested, source, list_level + 1)
        return
    if tag == 'pre':
        for line in ''.join(node.itertext()).rstrip('\n').split('\n'):
            p = doc.add_paragraph(style='Code')
            p.add_run(line or ' ')
        return
    if tag == 'table':
        table_block(doc, node, source)
        return
    if tag == 'hr':
        return
    for ch in node:
        block(doc, ch, source, list_level)

def diagram_detail_pages(doc):
    """Enlarge each half of the existing board without editing its source image."""
    section = doc.add_section(WD_SECTION_START.NEW_PAGE)
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width = Inches(11.7); section.page_height = Inches(8.3)
    for i, title in enumerate(('Application architecture diagram detail', 'Application workflow diagram detail')):
        if i:
            doc.add_page_break()
        doc.add_heading(title, 1)
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        shape = p.add_run().add_picture(str(ROOT / 'docs/diagrams/tldraw-review.png'), width=Inches(9.5))
        # srcRect controls the document view only; the original PNG stays intact.
        blipfill = shape._inline.xpath('.//pic:blipFill')[0]
        crop = OxmlElement('a:srcRect'); crop.set('t' if i else 'b', '50000')
        blipfill.insert(1, crop)
        shape._inline.extent.cy //= 2
        shape._inline.xpath('.//a:xfrm/a:ext')[0].set('cy', str(shape._inline.extent.cy))
        shape._inline.docPr.set('descr', title + ' from the original tldraw board')
        link(doc.add_paragraph(), 'Editable tldraw board', 'https://www.tldraw.com/f/MJ8ymcLzU-nKT2TWyaEVd', ROOT/'docs/APPLICATION_ARCHITECTURE.md')

for entry in json.loads(args.parsed.read_text(encoding='utf8')):
    source = ROOT / entry['source']
    if args.only and source.name != args.only:
        continue
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Inches(8.3); sec.page_height = Inches(11.7)
    sec.top_margin = sec.bottom_margin = Inches(.7)
    sec.left_margin = sec.right_margin = Inches(.75)
    sec.header_distance = sec.footer_distance = Inches(.3)
    styles = doc.styles
    for name in ('Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3', 'List Bullet', 'List Number'):
        styles[name].font.name = 'Calibri'
        styles[name].font.color.rgb = RGBColor(0, 0, 0)
    styles['Normal'].font.size = Pt(10.5)
    styles['Normal'].paragraph_format.space_after = Pt(7)
    styles['Normal'].paragraph_format.line_spacing = 1.1
    styles['Title'].font.size = Pt(25)
    styles['Title'].paragraph_format.space_after = Pt(13)
    for name, size in [('Heading 1', 16), ('Heading 2', 12.5), ('Heading 3', 11)]:
        styles[name].font.size = Pt(size)
        styles[name].paragraph_format.space_before = Pt(13)
        styles[name].paragraph_format.space_after = Pt(6)
    code = styles.add_style('Code', 1)
    code.font.name = 'Consolas'; code.font.size = Pt(8.5)
    code.paragraph_format.space_after = Pt(0)
    code.paragraph_format.line_spacing = 1
    code.paragraph_format.widow_control = False
    hp = sec.header.paragraphs[0]
    hp.add_run('Noteable  |  Application review collection').font.size = Pt(8)
    fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); fp._p.append(field)
    tree = html.fragment_fromstring(entry['html'], create_parent='div')
    if len(tree) == 0 or tree[0].tag != 'h1':
        doc.add_paragraph(clean_heading(source.stem.replace('_', ' ')), 'Title')
    first = True
    for node in tree:
        block(doc, node, source)
        if first:
            p = doc.add_paragraph()
            label = 'Historical document' if '/archive/' in entry['source'] else 'Review document'
            p.add_run(label + '  |  Source ').italic = True
            link(p, entry['source'], BASE + quote(entry['source']), source)
            first = False
    if source.name in ('APP_WORKFLOW.md', 'APPLICATION_ARCHITECTURE.md'):
        diagram_detail_pages(doc)
    if source.name == 'FEATURE_STATUS.md':
        # Heading spacing already separates these tables. Remove redundant
        # spacer paragraphs so the related links do not form a one-line page.
        for paragraph in list(doc.paragraphs):
            previous = paragraph._p.getprevious()
            if not paragraph.text and not paragraph._p.xpath('.//w:drawing') and previous is not None and previous.tag == qn('w:tbl'):
                paragraph._p.getparent().remove(paragraph._p)
    sub = 'archive' if '/archive/' in entry['source'] else 'current'
    target = OUT / sub / (source.stem + '.docx')
    target.parent.mkdir(parents=True, exist_ok=True)
    doc.core_properties.title = clean_heading(''.join(tree.xpath('./h1')[0].itertext())) if tree.xpath('./h1') else clean_heading(source.stem)
    doc.core_properties.subject = 'Noteable application review collection'
    doc.core_properties.author = 'Noteable project'
    doc.core_properties.keywords = entry['source']
    # Remove the template's inherited decorative rules, including Title residue.
    for border in list(doc.styles.element.xpath('.//w:pBdr')) + list(doc.element.xpath('.//w:pBdr')):
        border.getparent().remove(border)
    doc.save(target)
    record = {'source': entry['source'], 'source_sha256': hashlib.sha256(source.read_bytes()).hexdigest(), 'document': target.relative_to(ROOT).as_posix(), 'document_sha256': hashlib.sha256(target.read_bytes()).hexdigest(), 'embedded_images': len(doc.inline_shapes)}
    if args.only and any(x['source'] == entry['source'] for x in records):
        records = [record if x['source'] == entry['source'] else x for x in records]
    else:
        records.append(record)
manifest = ROOT / 'docs/publication/document-manifest.json'
manifest.write_text(json.dumps(records, indent=2) + '\n', encoding='utf8')
print(json.dumps({'documents': len(records), 'embedded_images': sum(x['embedded_images'] for x in records), 'manifest': str(manifest)}, indent=2))
